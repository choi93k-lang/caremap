"""
다운로드 폴더에서 공공데이터 원본을 읽어 울산광역시 데이터만 추출하는 스크립트입니다.
초보자도 쉽게 읽을 수 있도록 각 함수별로 역할을 단순하게 분리하였습니다.
"""

import os
import glob
import zipfile
import openpyxl
import pandas as pd


def extract_ulsan_hospitals(downloads_path, output_csv_path):
    """
    건강보험심사평가원 전국 병의원 ZIP 파일에서 울산광역시 병의원 데이터만 추출하여 CSV로 저장합니다.
    """
    # 1. 다운로드 폴더에서 가장 최근의 병원 ZIP 파일 찾기
    zip_files = sorted(glob.glob(os.path.join(downloads_path, "*.zip")), key=os.path.getmtime)
    target_zip = None
    for z_path in reversed(zip_files):
        try:
            with zipfile.ZipFile(z_path) as zf:
                names = zf.namelist()
                if any("병원" in name or "1." in name for name in names):
                    target_zip = z_path
                    break
        except Exception:
            continue

    if not target_zip:
        print("[오류] 다운로드 폴더에서 병의원 ZIP 파일을 찾을 수 없습니다.")
        return 0

    print(f"[1/2] 병의원 ZIP 파일 처리 시작: {os.path.basename(target_zip)}")

    # 2. ZIP 파일 내부의 '1.병원정보서비스' 엑셀 파일 열기
    with zipfile.ZipFile(target_zip) as zf:
        excel_name = [n for n in zf.namelist() if n.endswith(".xlsx") and ("1." in n or "병원" in n)][0]
        with zf.open(excel_name) as excel_file:
            workbook = openpyxl.load_workbook(excel_file, read_only=True)
            sheet = workbook.active
            
            # 헤더(첫 줄 컬럼 이름) 가져오기
            headers = [cell for cell in next(sheet.iter_rows(values_only=True))]
            sido_index = headers.index("시도코드명")

            # 3. 울산광역시(시도코드명: 울산) 행만 추출
            ulsan_rows = []
            for row in sheet.iter_rows(values_only=True):
                sido_value = str(row[sido_index]).strip()
                if sido_value == "울산":
                    ulsan_rows.append(row)

    # 4. 판다스 데이터프레임으로 변환 후 CSV로 저장
    df_hospitals = pd.DataFrame(ulsan_rows, columns=headers)
    df_hospitals.to_csv(output_csv_path, index=False, encoding="utf-8-sig")
    print(f"[완료] 울산 병의원 {len(df_hospitals)}개 추출 완료 -> {output_csv_path}")
    return len(df_hospitals)


def extract_ulsan_senior_centers(downloads_path, output_csv_path):
    """
    전국마을회관및경로당표준데이터 CSV 파일에서 울산광역시 경로당 데이터만 추출하여 CSV로 저장합니다.
    """
    # 1. 다운로드 폴더에서 가장 최근의 경로당 표준데이터 CSV 파일 찾기
    csv_candidates = glob.glob(os.path.join(downloads_path, "*마을회관*.csv")) + \
                     glob.glob(os.path.join(downloads_path, "*경로당*.csv")) + \
                     glob.glob(os.path.join(downloads_path, "*표준*.csv"))
    
    if not csv_candidates:
        print("[오류] 다운로드 폴더에서 경로당 CSV 파일을 찾을 수 없습니다.")
        return 0

    target_csv = sorted(csv_candidates, key=os.path.getmtime)[-1]
    print(f"[2/2] 경로당 CSV 파일 처리 시작: {os.path.basename(target_csv)}")

    # 2. CSV 파일 읽기 (인코딩: CP949)
    df_raw = pd.read_csv(target_csv, encoding="cp949", encoding_errors="replace")

    # 3. 울산광역시 데이터 필터링
    road_col = "소재지도로명주소" if "소재지도로명주소" in df_raw.columns else df_raw.columns[2]
    jibun_col = "소재지지번주소" if "소재지지번주소" in df_raw.columns else df_raw.columns[3]

    is_ulsan = df_raw[road_col].astype(str).str.contains("울산", na=False) | \
               df_raw[jibun_col].astype(str).str.contains("울산", na=False)
    df_ulsan = df_raw[is_ulsan].copy()

    # 4. 저장
    df_ulsan.to_csv(output_csv_path, index=False, encoding="utf-8-sig")
    print(f"[완료] 울산 경로당 {len(df_ulsan)}개 추출 완료 -> {output_csv_path}")
    return len(df_ulsan)


def run_extraction():
    """울산 공공데이터 추출 전체 실행 함수"""
    downloads_dir = os.path.expanduser("~/Downloads")
    output_dir = os.path.join(os.path.dirname(__file__), "..", "data", "raw")
    os.makedirs(output_dir, exist_ok=True)

    hospital_output_csv = os.path.join(output_dir, "ulsan_medical_official.csv")
    senior_output_csv = os.path.join(output_dir, "ulsan_senior_official.csv")

    hospital_count = extract_ulsan_hospitals(downloads_dir, hospital_output_csv)
    senior_count = extract_ulsan_senior_centers(downloads_dir, senior_output_csv)

    print(f"\n[최종 요약] 병의원: {hospital_count}개, 경로당: {senior_count}개 추출 완료!")


if __name__ == "__main__":
    run_extraction()
